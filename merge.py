import os
import glob
from copy import copy
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

dest_dir = "/Users/dognnguyen/Sites/keyword-searcher/destination"
output_path = "/Users/dognnguyen/Sites/keyword-searcher/reports/Keyword Planner - Tổng hợp.xlsx"

xlsx_files = sorted(glob.glob(os.path.join(dest_dir, "*.xlsx")))

# === Tạo workbook mới ===
merged_wb = Workbook()
# Xoá sheet mặc định
merged_wb.remove(merged_wb.active)

for xlsx_file in xlsx_files:
    # Lấy tên tab từ tên file (bỏ " - keyword")
    base_name = os.path.basename(xlsx_file).replace('.xlsx', '').replace(' - keyword', '').replace(' - keyword(1)', '')
    # Excel giới hạn tên tab 31 ký tự
    tab_name = base_name[:31]

    try:
        src_wb = load_workbook(xlsx_file)
        src_ws = src_wb.active

        # Tạo sheet mới trong file gộp
        new_ws = merged_wb.create_sheet(title=tab_name)

        # Copy từng cell: giá trị + style
        for row in src_ws.iter_rows():
            for cell in row:
                new_cell = new_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                # Copy style
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = cell.number_format

        # Copy merged cells
        for merged_range in src_ws.merged_cells.ranges:
            new_ws.merge_cells(str(merged_range))

        # Copy column widths
        for col_letter, col_dim in src_ws.column_dimensions.items():
            new_ws.column_dimensions[col_letter].width = col_dim.width

        # Copy row heights
        for row_num, row_dim in src_ws.row_dimensions.items():
            if row_dim.height:
                new_ws.row_dimensions[row_num].height = row_dim.height

        src_wb.close()
        print(f"✅ {tab_name}")

    except Exception as e:
        print(f"❌ {tab_name}: {e}")

# === Lưu file tổng hợp ===
merged_wb.save(output_path)
file_size_mb = os.path.getsize(output_path) / (1024 * 1024)
print(f"\n🎉 DONE! Saved to: {output_path}")
print(f"📦 File size: {file_size_mb:.2f} MB")
print(f"📋 Total tabs: {len(merged_wb.sheetnames)}")
print(f"📋 Tab names: {merged_wb.sheetnames}")
